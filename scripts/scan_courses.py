#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
酷学院课程批量上传 - 文件夹扫描脚本
功能：扫描指定文件夹，识别视频文件，生成上传计划Excel
"""

import os
import re
import argparse
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

# 视频文件扩展名
VIDEO_EXTENSIONS = {'.mp4', '.wmv', '.avi', '.mov', '.flv', '.mkv', '.webm'}
# PPT文件扩展名
PPT_EXTENSIONS = {'.pptx', '.ppt', '.pdf'}

# 资源分类映射规则
CATEGORY_MAPPING = {
    '销售族': '销售类',
    '职能族': '职能类',
    '运营族': '运营类',
    '技术族': '产研类',
    '产品研发': '产研类',
    '产品族': '产研类',
    '市场族': '运营类',
    '通用课程': '通识类',
    'M序列训练营': '管理类',
    '栋梁计划': '管理类',
    '英才计划': '管理类',
    '储干': '管理类',
    '主管论坛': '管理类',
    '新人交流会': '通识类',
}

# 识别为学习项目的关键词（同一专项下的课程合并）
PROJECT_KEYWORDS = [
    'M序列训练营', '栋梁计划', '英才计划', '储干',
    '主管论坛', '新人交流会', '内训师', '招聘课程',
    '人力赋能', '酷学院操作'
]


def is_video_file(filename):
    """判断是否为视频文件"""
    ext = Path(filename).suffix.lower()
    return ext in VIDEO_EXTENSIONS


def is_ppt_file(filename):
    """判断是否为PPT/PDF文件"""
    ext = Path(filename).suffix.lower()
    return ext in PPT_EXTENSIONS


def extract_date_from_name(filename):
    """从文件名提取六位数日期"""
    # 优先匹配完整的8位日期（如20230614）
    match = re.search(r'(?<!\d)(\d{8})(?!\d)', filename)
    if match:
        date_str = match.group(1)
        # 20230604 -> 230604
        return date_str[2:]
    
    # 匹配中文日期格式
    match = re.search(r'(\d{4})年(\d{1,2})月(\d{1,2})日', filename)
    if match:
        year = match.group(1)[2:]  # 2023 -> 23
        month = match.group(2).zfill(2)
        day = match.group(3).zfill(2)
        return f"{year}{month}{day}"
    
    # 匹配横杠日期格式
    match = re.search(r'(\d{4})-(\d{1,2})-(\d{1,2})', filename)
    if match:
        year = match.group(1)[2:]
        month = match.group(2).zfill(2)
        day = match.group(3).zfill(2)
        return f"{year}{month}{day}"
    
    # 最后匹配独立的6位数字日期（前后不能是数字，避免从8位中截取）
    match = re.search(r'(?<!\d)(\d{6})(?!\d)', filename)
    if match:
        return match.group(1)
    
    return None


def clean_course_name(filename):
    """从文件名提取课程名称"""
    name = Path(filename).stem
    # 去掉常见的前缀和后缀
    name = re.sub(r'^\d{8}-', '', name)  # 去掉时间戳前缀
    name = re.sub(r'_video$', '', name)
    name = re.sub(r'_audio$', '', name)
    # 去掉日期部分（先去掉长的，再去掉短的，避免截断）
    name = re.sub(r'\s*\d{8}\s*$', '', name)  # 先去8位日期
    name = re.sub(r'\s*\d{6}\s*$', '', name)  # 再去6位日期
    name = re.sub(r'\s*\d{4}年\d{1,2}月\d{1,2}日\s*$', '', name)
    name = re.sub(r'\s*\d{4}-\d{1,2}-\d{1,2}\s*$', '', name)  # 横杠日期
    return name.strip()


def match_category(folder_path, root_folder):
    """根据文件夹路径匹配资源分类"""
    # 先检查根文件夹名称
    root_name = os.path.basename(root_folder)
    for keyword, category in CATEGORY_MAPPING.items():
        if keyword in root_name:
            return category
    
    # 再检查相对路径的各部分
    rel_path = os.path.relpath(folder_path, root_folder)
    parts = rel_path.split(os.sep)
    
    for part in parts:
        for keyword, category in CATEGORY_MAPPING.items():
            if keyword in part:
                return category
    
    return '通识类'  # 默认分类


def identify_project(folder_name):
    """判断文件夹是否为学习项目"""
    for keyword in PROJECT_KEYWORDS:
        if keyword in folder_name:
            return True
    return False


def scan_folder(root_folder):
    """扫描文件夹，返回课程列表"""
    courses = []
    
    for root, dirs, files in os.walk(root_folder):
        # 过滤隐藏文件夹和临时文件夹
        dirs[:] = [d for d in dirs if not d.startswith('.') and not d.startswith('~$')]
        
        video_files = [f for f in files if is_video_file(f)]
        ppt_files = [f for f in files if is_ppt_file(f)]
        
        if not video_files:
            continue
        
        # 获取相对路径和分类
        rel_path = os.path.relpath(root, root_folder)
        category = match_category(root, root_folder)
        is_project = identify_project(os.path.basename(root))
        
        # 为每个视频创建课程条目
        for video in video_files:
            video_path = os.path.join(root, video)
            course_name = clean_course_name(video)
            date_str = extract_date_from_name(video)
            
            # 匹配同名PPT
            base_name = Path(video).stem
            related_ppts = []
            for ppt in ppt_files:
                ppt_base = Path(ppt).stem
                # 模糊匹配：去掉日期后比较
                if clean_course_name(ppt_base) in clean_course_name(video) or \
                   clean_course_name(video) in clean_course_name(ppt_base):
                    related_ppts.append(ppt)
            
            # 生成完整课程名称
            if date_str:
                full_name = f"{course_name} {date_str}"
            else:
                full_name = course_name
            
            courses.append({
                'folder': rel_path,
                'video_file': video,
                'video_path': video_path,
                'course_name': full_name,
                'category': category,
                'is_project_folder': is_project,
                'related_ppts': related_ppts,
                'file_size_mb': round(os.path.getsize(video_path) / 1024 / 1024, 1)
            })
    
    return courses


def generate_excel(courses, output_path):
    """生成上传计划Excel"""
    wb = Workbook()
    ws = wb.active
    ws.title = "上传计划"
    
    # 表头
    headers = ['序号', '课程名称', '资源分类', '所属文件夹', '视频文件', '配套PPT', '文件大小(MB)', '是否项目文件夹', '操作']
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
    
    # 数据行
    for i, course in enumerate(courses, 2):
        ws.cell(row=i, column=1, value=i-1)
        ws.cell(row=i, column=2, value=course['course_name'])
        ws.cell(row=i, column=3, value=course['category'])
        ws.cell(row=i, column=4, value=course['folder'])
        ws.cell(row=i, column=5, value=course['video_file'])
        ws.cell(row=i, column=6, value=', '.join(course['related_ppts']) if course['related_ppts'] else '')
        ws.cell(row=i, column=7, value=course['file_size_mb'])
        ws.cell(row=i, column=8, value='是' if course['is_project_folder'] else '否')
        ws.cell(row=i, column=9, value='待上传')
    
    # 自动调整列宽
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column].width = adjusted_width
    
    # 添加统计sheet
    ws2 = wb.create_sheet("统计信息")
    ws2.cell(row=1, column=1, value="课程总数").font = Font(bold=True)
    ws2.cell(row=1, column=2, value=len(courses))
    
    # 按分类统计
    category_count = {}
    for c in courses:
        cat = c['category']
        category_count[cat] = category_count.get(cat, 0) + 1
    
    ws2.cell(row=3, column=1, value="分类统计").font = Font(bold=True)
    row = 4
    for cat, count in sorted(category_count.items()):
        ws2.cell(row=row, column=1, value=cat)
        ws2.cell(row=row, column=2, value=count)
        row += 1
    
    wb.save(output_path)
    print(f"✅ 上传计划已生成: {output_path}")
    print(f"   共 {len(courses)} 个课程")
    for cat, count in sorted(category_count.items()):
        print(f"   {cat}: {count}个")


def main():
    parser = argparse.ArgumentParser(description='酷学院课程批量上传 - 文件夹扫描工具')
    parser.add_argument('--folder', required=True, help='要扫描的文件夹路径')
    parser.add_argument('--output', default='上传计划.xlsx', help='输出Excel文件路径')
    
    args = parser.parse_args()
    
    root_folder = args.folder
    if not os.path.exists(root_folder):
        print(f"❌ 文件夹不存在: {root_folder}")
        return
    
    print(f"🔍 正在扫描文件夹: {root_folder}")
    courses = scan_folder(root_folder)
    
    if not courses:
        print("❌ 未找到视频文件")
        return
    
    print(f"\n📊 找到 {len(courses)} 个视频课程")
    generate_excel(courses, args.output)


if __name__ == '__main__':
    main()
